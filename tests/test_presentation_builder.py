from pptx import Presentation

from src.utils.config import PROJECT_ROOT

from scripts.shared.reporting.build_presentation import CHARTS, DEFAULT_OUT, \
    build_deck

EXPECTED_SLIDES = 11
EXPECTED_CHART_EMBEDS = 5  # 06 on results, 01+02 on deep dive, 04+03 on NEI


def _count_pictures(prs):
    return sum(
        getattr(shape, 'shape_type', None) == 13
        for slide in prs.slides for shape in slide.shapes
    )


def test_v5_deck_builds_with_expected_slides(tmp_path):
    out = build_deck(tmp_path / 'deck.pptx')

    assert out == tmp_path / 'deck.pptx'
    assert out.exists()
    prs = Presentation(str(out))
    assert len(prs.slides) == EXPECTED_SLIDES
    assert _count_pictures(prs) == EXPECTED_CHART_EMBEDS


def test_v5_deck_chart_assets_are_committed():
    names = [
        '01_per_class_metrics.png',
        '02_confusion_matrix.png',
        '03_risk_coverage.png',
        '04_nei_outcomes.png',
        '06_validation_vs_test.png',
    ]
    for name in names:
        assert (CHARTS / name).exists(), name


def test_v5_deck_default_output_path():
    assert DEFAULT_OUT == PROJECT_ROOT / 'presentation' / \
        'claim_level_hallucination_detection.pptx'
